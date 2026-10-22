from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = {
    'nome_empresa': ('Empresa / cliente', 45),
    'cpf_cnpj': ('CPF / CNPJ', 22),
    'tipo': ('Tipo de pessoa', 18),
    'responsavel': ('Responsável', 35),
    'data_vencimento': ('Vencimento', 18),
    'dias': ('Dias para vencer', 20),
    'status': ('Situação', 23),
    'email_contato': ('E-mail', 38),
    'telefone_contato': ('Telefone', 22),
    'arquivo_drive_id': ('Arquivo do certificado', 48),
}


def gerar_excel_certificados(certificados, colunas):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Certificados'
    ws.append([COLUNAS[c][0] for c in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor='1E3A5F')
        celula.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28
    for certificado in certificados:
        valores = {c: getattr(certificado, c, '') for c in colunas}
        valores['dias'] = certificado.dias_para_vencimento()
        valores['status'] = certificado.status_vencimento()
        ws.append([valores[c] for c in colunas])
        for indice, coluna in enumerate(colunas, 1):
            celula = ws.cell(ws.max_row, indice)
            # Dados cadastrais são texto, nunca fórmulas do Excel.
            if isinstance(celula.value, str):
                celula.data_type = 's'
            celula.alignment = Alignment(vertical='top', wrap_text=True)
            if coluna == 'data_vencimento':
                celula.number_format = 'dd/mm/yyyy'
            elif coluna in {'cpf_cnpj', 'telefone_contato'}:
                celula.number_format = '@'
            if ws.max_row % 2 == 0:
                celula.fill = PatternFill('solid', fgColor='F1F5F9')
    for i, coluna in enumerate(colunas, 1):
        ws.column_dimensions[get_column_letter(i)].width = COLUNAS[coluna][1]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)
    return saida
