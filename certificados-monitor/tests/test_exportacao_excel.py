import io
import unittest
from datetime import date

from flask import Flask
from openpyxl import load_workbook
from src.models.certificado import Certificado, db
from src.routes.certificado import certificado_bp


class ExportacaoExcelTestCase(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI='sqlite:///:memory:', SQLALCHEMY_TRACK_MODIFICATIONS=False)
        db.init_app(self.app)
        self.app.register_blueprint(certificado_bp, url_prefix='/api')
        with self.app.app_context():
            db.create_all()
            db.session.add_all([
                Certificado(nome_empresa='=HYPERLINK("https://example.com")', cpf_cnpj='00123456000190', tipo='PJ', data_vencimento=date(2027, 1, 15)),
                Certificado(nome_empresa='Outra empresa', cpf_cnpj='22123456000190', tipo='PJ', data_vencimento=date(2026, 1, 15)),
                Certificado(nome_empresa='Inativo', cpf_cnpj='33123456000190', tipo='PJ', data_vencimento=date(2026, 1, 15), ativo=False),
            ])
            db.session.commit()
        self.client = self.app.test_client()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def exportar(self, query):
        return self.client.get('/api/certificados/exportar-excel?' + query)

    def test_exporta_colunas_selecionadas_preservando_documento_e_sem_formulas(self):
        r = self.exportar('colunas=nome_empresa,cpf_cnpj,data_vencimento')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.content_type)
        self.assertIn('attachment', r.headers['Content-Disposition'])
        ws = load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.max_column, 3)
        self.assertEqual(ws['B2'].value, '00123456000190')
        self.assertEqual(ws['A2'].data_type, 's')
        self.assertEqual(ws['C2'].number_format, 'dd/mm/yyyy')
        self.assertEqual(ws.freeze_panes, 'A2')

    def test_filtro_e_ordem_independentes_da_paginacao(self):
        ws = load_workbook(io.BytesIO(self.exportar('colunas=cpf_cnpj&ids=2,1').data)).active
        self.assertEqual(ws['A2'].value, '22123456000190')
        self.assertEqual(ws['A3'].value, '00123456000190')
        ws = load_workbook(io.BytesIO(self.exportar('colunas=cpf_cnpj&ids=1').data)).active
        self.assertEqual(ws.max_row, 2)

    def test_rejeita_selecoes_invalidas(self):
        for query in ['colunas=', 'colunas=senha', 'colunas=cpf_cnpj,cpf_cnpj', 'colunas=cpf_cnpj&ids=', 'colunas=cpf_cnpj&ids=xyz', 'colunas=cpf_cnpj&ids=-1', 'colunas=cpf_cnpj&ids=999', 'colunas=cpf_cnpj&ids=3']:
            with self.subTest(query=query):
                self.assertEqual(self.exportar(query).status_code, 400)
